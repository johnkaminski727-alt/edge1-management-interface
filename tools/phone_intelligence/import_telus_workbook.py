#!/usr/bin/env python3

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_SHA256 = (
    "25694b726885e3f2a4417623b76e2dc7"
    "ba79b5a9dda386acdd0dac437ea0bdef"
)

EXPECTED_NUMBERS = 677
EXPECTED_OCCURRENCES = 4255
EXPECTED_BILL_REFERENCES = 1566
EXPECTED_SOURCE_FILES = 50

DIGITS = re.compile(r"\D+")


def utc_now():
    return datetime.now(timezone.utc).isoformat(
        timespec="seconds"
    ).replace("+00:00", "Z")


def clean(value):
    if value is None:
        return None

    value = str(value).strip()
    return value or None


def integer(value):
    if value is None or value == "":
        return 0

    return int(value)


def sha256_file(path):
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def normalize_phone(value):
    raw = clean(value)

    if not raw:
        raise ValueError("empty phone number")

    digits = DIGITS.sub("", raw)

    if len(digits) == 10:
        return "+1" + digits, digits

    if len(digits) == 11 and digits.startswith("1"):
        return "+" + digits, digits[1:]

    if raw.startswith("+") and 7 <= len(digits) <= 15:
        return "+" + digits, digits

    raise ValueError(f"unsupported telephone number: {raw!r}")


def normalized_status(value):
    original = (clean(value) or "").lower()

    if original == "confirmed":
        return "confirmed"

    if original == "probable":
        return "probable"

    return "unresolved"


def source_files(value):
    raw = clean(value)

    if not raw:
        return []

    text = raw.replace("\r", "\n")

    results = []

    for line in text.split("\n"):
        for semicolon in line.split(";"):
            for comma in semicolon.split(","):
                item = comma.strip()

                if item:
                    results.append(item)

    return list(dict.fromkeys(results))


def load_records(path):
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook["ALL NUMBERS"]

        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator)

        records = []

        for row_number, row in enumerate(iterator, start=2):
            if not any(
                value is not None and str(value).strip()
                for value in row
            ):
                continue

            record = dict(zip(headers, row))

            if not clean(record.get("PHONE NUMBER")):
                continue

            record["_row_number"] = row_number
            records.append(record)

        return records

    finally:
        workbook.close()


def analyze(records):
    numbers = set()
    occurrence_total = 0
    bill_reference_total = 0
    sources = set()

    for record in records:
        normalized, _ = normalize_phone(
            record["PHONE NUMBER"]
        )

        if normalized in numbers:
            raise ValueError(
                f"duplicate normalized telephone number: {normalized}"
            )

        numbers.add(normalized)

        occurrence_total += integer(
            record.get("OCCURRENCES")
        )

        bill_reference_total += integer(
            record.get("BILLS")
        )

        sources.update(
            source_files(record.get("SOURCE BILL FILES"))
        )

    return {
        "numbers": len(numbers),
        "occurrences": occurrence_total,
        "bill_references": bill_reference_total,
        "source_files": len(sources),
    }


def verify_baseline(path, records):
    digest = sha256_file(path)

    if digest != EXPECTED_SHA256:
        raise RuntimeError(
            "workbook SHA256 mismatch:\n"
            f" expected: {EXPECTED_SHA256}\n"
            f" actual:   {digest}"
        )

    result = analyze(records)

    expected = {
        "numbers": EXPECTED_NUMBERS,
        "occurrences": EXPECTED_OCCURRENCES,
        "bill_references": EXPECTED_BILL_REFERENCES,
        "source_files": EXPECTED_SOURCE_FILES,
    }

    if result != expected:
        raise RuntimeError(
            "workbook reconciliation failure:\n"
            f" expected={expected}\n"
            f" actual={result}"
        )

    return digest, result


def connect_database(path):
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row

    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA busy_timeout=5000")

    return conn


def ensure_empty(conn):
    tables = (
        "phone_numbers",
        "organizations",
        "associations",
        "source_documents",
        "occurrences",
        "evidence",
        "import_batches",
        "audit_events",
    )

    counts = {}

    for table in tables:
        counts[table] = conn.execute(
            f'SELECT COUNT(*) FROM "{table}"'
        ).fetchone()[0]

    if any(counts.values()):
        raise RuntimeError(
            "database is not empty; refusing initial import: "
            + json.dumps(counts, sort_keys=True)
        )


def import_records(conn, workbook_path, digest, records):
    now = utc_now()

    association_count = 0
    organization_count = 0
    source_count = 0

    organizations = {}
    sources = {}

    for record in records:
        display = clean(record["PHONE NUMBER"])

        normalized, national = normalize_phone(display)

        status = normalized_status(
            record.get("EVIDENCE STATUS")
        )

        npa = national[:3] if len(national) == 10 else None

        occurrence_count = integer(
            record.get("OCCURRENCES")
        )

        conn.execute(
            """
            INSERT INTO phone_numbers(
                normalized_number,
                display_number,
                country_code,
                national_number,
                npa,
                status,
                occurrence_count,
                created_at,
                updated_at
            )
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (
                normalized,
                display,
                "1" if normalized.startswith("+1") else None,
                national,
                npa,
                status,
                occurrence_count,
                now,
                now,
            ),
        )

        phone_id = conn.execute(
            """
            SELECT id
            FROM phone_numbers
            WHERE normalized_number=?
            """,
            (normalized,),
        ).fetchone()["id"]

        association_name = clean(
            record.get("PUBLIC ASSOCIATION")
        )

        evidence_status = clean(
            record.get("EVIDENCE STATUS")
        )

        research_note = clean(
            record.get("RESEARCH NOTE")
        )

        location = clean(record.get("LOCATION"))
        category = clean(record.get("CATEGORY"))
        source_url = clean(record.get("SOURCE URL"))
        source_basis = clean(record.get("SOURCE BASIS"))

        organization_id = None

        if association_name:
            org_key = (
                association_name,
                location,
                category,
            )

            if org_key not in organizations:
                conn.execute(
                    """
                    INSERT INTO organizations(
                        name,
                        organization_type,
                        region,
                        website,
                        notes,
                        created_at,
                        updated_at
                    )
                    VALUES(?,?,?,?,?,?,?)
                    """,
                    (
                        association_name,
                        category,
                        location,
                        source_url,
                        None,
                        now,
                        now,
                    ),
                )

                organization_id = conn.execute(
                    "SELECT last_insert_rowid()"
                ).fetchone()[0]

                organizations[org_key] = organization_id
                organization_count += 1

            else:
                organization_id = organizations[org_key]

        conn.execute(
            """
            INSERT INTO associations(
                phone_number_id,
                organization_id,
                association_label,
                confidence,
                research_notes,
                created_at,
                updated_at
            )
            VALUES(?,?,?,?,?,?,?)
            """,
            (
                phone_id,
                organization_id,
                association_name,
                status,
                research_note,
                now,
                now,
            ),
        )

        association_id = conn.execute(
            "SELECT last_insert_rowid()"
        ).fetchone()[0]

        association_count += 1

        if source_url or source_basis or evidence_status:
            conn.execute(
                """
                INSERT INTO evidence(
                    association_id,
                    evidence_type,
                    source_title,
                    source_url,
                    source_reference,
                    evidence_summary,
                    verification_status,
                    retrieved_at,
                    created_at
                )
                VALUES(?,?,?,?,?,?,?,?,?)
                """,
                (
                    association_id,
                    "public-research",
                    source_basis,
                    source_url,
                    clean(record.get("ARCHIVE")),
                    (
                        f"Original workbook evidence status: "
                        f"{evidence_status}"
                    ),
                    (
                        "verified"
                        if status == "confirmed"
                        else "unverified"
                    ),
                    None,
                    now,
                ),
            )

        for filename in source_files(
            record.get("SOURCE BILL FILES")
        ):
            if filename not in sources:
                conn.execute(
                    """
                    INSERT INTO source_documents(
                        document_type,
                        source_name,
                        source_reference,
                        verification_status,
                        notes,
                        created_at
                    )
                    VALUES(?,?,?,?,?,?)
                    """,
                    (
                        "telus-bill",
                        filename,
                        filename,
                        "missing",
                        (
                            "Referenced by imported phone "
                            "association workbook; underlying "
                            "source document was not independently "
                            "recovered during this audit."
                        ),
                        now,
                    ),
                )

                source_id = conn.execute(
                    "SELECT last_insert_rowid()"
                ).fetchone()[0]

                sources[filename] = source_id
                source_count += 1

            conn.execute(
                """
                INSERT INTO occurrences(
                    phone_number_id,
                    source_document_id,
                    occurrence_type,
                    raw_value,
                    notes,
                    created_at
                )
                VALUES(?,?,?,?,?,?)
                """,
                (
                    phone_id,
                    sources[filename],
                    "bill-reference",
                    display,
                    (
                        "Workbook-level source relationship. "
                        "Aggregate occurrence count retained "
                        "on phone_numbers."
                    ),
                    now,
                ),
            )

    conn.execute(
        """
        INSERT INTO import_batches(
            import_name,
            source_filename,
            source_sha256,
            row_count,
            imported_count,
            skipped_count,
            imported_at,
            notes
        )
        VALUES(?,?,?,?,?,?,?,?)
        """,
        (
            "TELUS Phone Association Master",
            workbook_path.name,
            digest,
            len(records),
            len(records),
            0,
            now,
            (
                "Initial authoritative import. "
                "Workbook aggregate occurrence count=4255; "
                "bill-reference total=1566; "
                "unique source bills=50."
            ),
        ),
    )

    conn.execute(
        """
        INSERT INTO audit_events(
            event_type,
            entity_type,
            entity_id,
            actor,
            detail_json,
            created_at
        )
        VALUES(?,?,?,?,?,?)
        """,
        (
            "import.completed",
            "dataset",
            digest,
            "phone-intelligence-importer",
            json.dumps(
                {
                    "source": workbook_path.name,
                    "sha256": digest,
                    "phone_numbers": len(records),
                    "organizations": organization_count,
                    "associations": association_count,
                    "source_documents": source_count,
                    "aggregate_occurrences":
                        EXPECTED_OCCURRENCES,
                    "bill_references":
                        EXPECTED_BILL_REFERENCES,
                },
                sort_keys=True,
            ),
            now,
        ),
    )


def validate_database(conn):
    counts = {}

    for table in (
        "phone_numbers",
        "organizations",
        "associations",
        "source_documents",
        "occurrences",
        "evidence",
        "import_batches",
        "audit_events",
    ):
        counts[table] = conn.execute(
            f'SELECT COUNT(*) FROM "{table}"'
        ).fetchone()[0]

    aggregate_occurrences = conn.execute(
        """
        SELECT COALESCE(SUM(occurrence_count),0)
        FROM phone_numbers
        """
    ).fetchone()[0]

    statuses = dict(
        conn.execute(
            """
            SELECT status, COUNT(*)
            FROM phone_numbers
            GROUP BY status
            ORDER BY status
            """
        ).fetchall()
    )

    integrity = conn.execute(
        "PRAGMA integrity_check"
    ).fetchone()[0]

    foreign_key_errors = conn.execute(
        "PRAGMA foreign_key_check"
    ).fetchall()

    if counts["phone_numbers"] != EXPECTED_NUMBERS:
        raise RuntimeError("phone count reconciliation failed")

    if counts["associations"] != EXPECTED_NUMBERS:
        raise RuntimeError("association count reconciliation failed")

    if counts["source_documents"] != EXPECTED_SOURCE_FILES:
        raise RuntimeError("source-document reconciliation failed")

    if aggregate_occurrences != EXPECTED_OCCURRENCES:
        raise RuntimeError("occurrence reconciliation failed")

    if statuses != {
        "confirmed": 48,
        "probable": 2,
        "unresolved": 627,
    }:
        raise RuntimeError(
            f"status reconciliation failed: {statuses}"
        )

    if integrity != "ok":
        raise RuntimeError(
            f"SQLite integrity check failed: {integrity}"
        )

    if foreign_key_errors:
        raise RuntimeError(
            f"foreign-key violations: {foreign_key_errors}"
        )

    return counts, statuses, aggregate_occurrences


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--workbook",
        required=True,
        type=Path,
    )

    parser.add_argument(
        "--database",
        required=True,
        type=Path,
    )

    parser.add_argument(
        "--apply",
        action="store_true",
        help="Commit import. Without this option only validate.",
    )

    args = parser.parse_args()

    records = load_records(args.workbook)

    digest, analysis = verify_baseline(
        args.workbook,
        records,
    )

    print("Workbook reconciliation: PASS")
    print("SHA256:", digest)
    print("Numbers:", analysis["numbers"])
    print("Occurrences:", analysis["occurrences"])
    print("Bill references:", analysis["bill_references"])
    print("Source files:", analysis["source_files"])

    if not args.apply:
        print()
        print("DRY RUN: PASS")
        print("No database changes made.")
        return

    conn = connect_database(args.database)

    try:
        ensure_empty(conn)

        conn.execute("BEGIN IMMEDIATE")

        import_records(
            conn,
            args.workbook,
            digest,
            records,
        )

        counts, statuses, occurrences = validate_database(
            conn
        )

        conn.commit()

    except Exception:
        conn.rollback()
        raise

    finally:
        conn.close()

    print()
    print("IMPORT: COMMITTED")
    print("Counts:", json.dumps(counts, sort_keys=True))
    print("Statuses:", json.dumps(statuses, sort_keys=True))
    print("Aggregate occurrences:", occurrences)


if __name__ == "__main__":
    main()
